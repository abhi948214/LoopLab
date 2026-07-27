#!/usr/bin/env python3
"""Generate Princhan.pptx — LoopLAB project presentation."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BRAND_BLUE = RGBColor(0, 142, 214)
DARK = RGBColor(51, 51, 51)
WHITE = RGBColor(255, 255, 255)
LIGHT_GRAY = RGBColor(245, 245, 245)
MUTED = RGBColor(120, 120, 120)


def set_slide_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(3.1), prs.slide_width, Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BRAND_BLUE
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.4), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT

    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.4), Inches(8.4), Inches(1.2))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(22)
    p.font.color.rgb = BRAND_BLUE
    p.alignment = PP_ALIGN.LEFT

    footer = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(8.4), Inches(0.5))
    fp = footer.text_frame.paragraphs[0]
    fp.text = "LoopLAB · Social Media Platform"
    fp.font.size = Pt(14)
    fp.font.color.rgb = MUTED


def add_section_slide(prs, section_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BRAND_BLUE)

    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(8.4), Inches(1.5))
    p = box.text_frame.paragraphs[0]
    p.text = section_title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT


def add_bullet_slide(prs, title, bullets, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), prs.slide_height)
    accent.fill.solid()
    accent.fill.fore_color.rgb = BRAND_BLUE
    accent.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(8.6), Inches(0.9))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(32)
    tp.font.bold = True
    tp.font.color.rgb = DARK

    body_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(8.2), Inches(4.8))
    tf = body_box.text_frame
    tf.word_wrap = True

    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(20)
        p.font.color.rgb = DARK
        p.space_after = Pt(14)

    if note:
        note_box = slide.shapes.add_textbox(Inches(0.9), Inches(6.2), Inches(8.2), Inches(0.6))
        np = note_box.text_frame.paragraphs[0]
        np.text = note
        np.font.size = Pt(14)
        np.font.italic = True
        np.font.color.rgb = MUTED


def add_two_column_slide(prs, title, left_title, left_items, right_title, right_items):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), prs.slide_height)
    accent.fill.solid()
    accent.fill.fore_color.rgb = BRAND_BLUE
    accent.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(8.6), Inches(0.9))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(32)
    tp.font.bold = True
    tp.font.color.rgb = DARK

    for col_x, col_title, items in [
        (Inches(0.8), left_title, left_items),
        (Inches(5.1), right_title, right_items),
    ]:
        head = slide.shapes.add_textbox(col_x, Inches(1.4), Inches(3.8), Inches(0.5))
        hp = head.text_frame.paragraphs[0]
        hp.text = col_title
        hp.font.size = Pt(20)
        hp.font.bold = True
        hp.font.color.rgb = BRAND_BLUE

        body = slide.shapes.add_textbox(col_x, Inches(1.95), Inches(3.8), Inches(4.5))
        tf = body.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = item
            p.font.size = Pt(17)
            p.font.color.rgb = DARK
            p.space_after = Pt(10)


def add_thank_you_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK)

    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(8.4), Inches(1.5))
    p = box.text_frame.paragraphs[0]
    p.text = "Thank You"
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    sub = slide.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(8.4), Inches(1))
    sp = sub.text_frame.paragraphs[0]
    sp.text = "Questions & Discussion"
    sp.font.size = Pt(22)
    sp.font.color.rgb = BRAND_BLUE
    sp.alignment = PP_ALIGN.CENTER


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Princhan",
        "LoopLAB — Build Social Profiles & Grow Revenue",
    )

    add_section_slide(prs, "Overview")

    add_bullet_slide(
        prs,
        "What is LoopLAB?",
        [
            "A social media platform landing page for creators and brands",
            "Helps users build profiles, create content, and share with audiences",
            "Designed for engagement, growth, and monetization",
            "Modern responsive UI built with Bootstrap 4",
        ],
    )

    add_bullet_slide(
        prs,
        "Problem Statement",
        [
            "Creators struggle to build a consistent online presence",
            "Content ideas, captions, and hashtags take too much time",
            "Hard to connect exploration, creation, and sharing in one flow",
            "Need a simple platform to grow audience and revenue",
        ],
    )

    add_section_slide(prs, "Platform Features")

    add_two_column_slide(
        prs,
        "Core Sections",
        "Explore",
        [
            "Discover trends and communities",
            "Connect with like-minded creators",
            "Find inspiration for new content",
        ],
        "Create",
        [
            "Turn passion into posts and reels",
            "Structured content creation workflow",
            "Build a recognizable personal brand",
        ],
    )

    add_bullet_slide(
        prs,
        "Share & Monetize",
        [
            "Publish content across social channels from one hub",
            "Grow followers and engagement over time",
            "Convert audience attention into revenue opportunities",
            "Sign-up flow for onboarding new users quickly",
        ],
        note="Home · Explore · Create · Share — the full creator journey",
    )

    add_section_slide(prs, "Instagram AI Assistant")

    add_bullet_slide(
        prs,
        "AI-Powered Content Tools",
        [
            "Caption Generator — tone and length tailored to your post",
            "Hashtag Generator — relevant tags for better reach",
            "Bio Generator — compelling profile bios in seconds",
            "Content Ideas — brainstorm posts and Reels for any niche",
        ],
        note="Powered by OpenAI GPT-4o-mini via Node.js + Express backend",
    )

    add_two_column_slide(
        prs,
        "Technology Stack",
        "Frontend",
        [
            "HTML5 + Bootstrap 4",
            "Custom CSS (brand color #008ed6)",
            "jQuery scrollspy & smooth scrolling",
            "Font Awesome icons",
        ],
        "Backend & AI",
        [
            "Node.js + Express API server",
            "OpenAI GPT-4o-mini integration",
            "REST endpoints for each AI tool",
            "Environment-based API key config",
        ],
    )

    add_bullet_slide(
        prs,
        "API Endpoints",
        [
            "GET  /api/health — server & AI status check",
            "POST /api/generate/caption — Instagram caption",
            "POST /api/generate/hashtags — hashtag suggestions",
            "POST /api/generate/bio — profile bio text",
            "POST /api/generate/ideas — content brainstorming",
        ],
    )

    add_section_slide(prs, "Roadmap")

    add_bullet_slide(
        prs,
        "Future Scope",
        [
            "User authentication and profile dashboards",
            "Analytics for post performance and audience growth",
            "Direct Instagram / social media integrations",
            "Scheduling and auto-publishing for creators",
            "Premium AI features and team collaboration",
        ],
    )

    add_thank_you_slide(prs)

    output = "/workspace/Princhan.pptx"
    prs.save(output)
    print(f"Created {output} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
